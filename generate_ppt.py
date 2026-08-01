from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "An Intelligent Cybercrime Reporting and Complaint Management System using Machine Learning and NLP for Analyzing and Classifying Cybercrime Complaints"
    
    # Make the heading smaller due to length
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(28)

    subtitle.text = "Presented By:\nNAME:USN\nTeam No:01\n\nUnder the Guidance of\nGuide Name\nAssistant Professor(Sr.)\nDept. of CSE\nSMVITM, Bantakal"

    # Problem Statement
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Problem Statement"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "Project Track"
    p.font.bold = True
    p.font.size = Pt(24)
    p = tf.add_paragraph()
    p.text = "• Applied Engineering-Industry-Based Projects"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• The rapid increase in cybercrime requires an automated, intelligent system to classify and manage complaints effectively, reducing manual effort and improving response times."
    p.level = 1

    # Objectives
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Objectives"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "• To develop a full-stack platform for reporting cybercrimes seamlessly."
    p = tf.add_paragraph()
    p.text = "• To utilize Machine Learning (TF-IDF + Logistic Regression) for automatic categorization of complaints into 11 distinct categories."
    p = tf.add_paragraph()
    p.text = "• To provide real-time updates and notifications using WebSockets."
    p = tf.add_paragraph()
    p.text = "• To ensure secure and scalable management of user data and evidence."

    # Introduction
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Introduction"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "• CyberGuard is an AI-powered cybercrime reporting platform."
    p = tf.add_paragraph()
    p.text = "• It features a web portal for users and administrators, and a mobile application for on-the-go reporting."
    p = tf.add_paragraph()
    p.text = "• The system integrates a Node.js backend, a React web frontend, a React Native mobile app, and a Python Flask ML service."

    # Methodology
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Methodology"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "• Data Collection & Preprocessing: NLP techniques (NLTK) applied to user complaint text."
    p = tf.add_paragraph()
    p.text = "• Model Training: TF-IDF vectorization paired with Logistic Regression for text classification."
    p = tf.add_paragraph()
    p.text = "• Backend Integration: Node.js/Express API interfacing with MongoDB for data persistence."
    p = tf.add_paragraph()
    p.text = "• Frontend & Mobile: React and Expo used for responsive and accessible user interfaces."

    # Results
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Results"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "• Successfully automated the categorization of 11 cybercrime types."
    p = tf.add_paragraph()
    p.text = "• Real-time synchronization achieved across web and mobile platforms using Socket.IO."
    p = tf.add_paragraph()
    p.text = "• Enhanced user experience with interactive 3D UI (Three.js) and data visualizations (Recharts)."

    # Conclusion
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Conclusion"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "• CyberGuard provides a robust, scalable, and intelligent solution for cybercrime complaint management."
    p = tf.add_paragraph()
    p.text = "• The integration of AI significantly streamlines the triage process for law enforcement."
    p = tf.add_paragraph()
    p.text = "• Future work may include deep learning models and multi-lingual support."

    # References
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "References"
    tf = content.text_frame
    p = tf.add_paragraph()
    p.text = "[1] React Documentation. (n.d.). Retrieved from reactjs.org"
    p = tf.add_paragraph()
    p.text = "[2] Node.js Documentation. (n.d.). Retrieved from nodejs.org"
    p = tf.add_paragraph()
    p.text = "[3] Scikit-learn: Machine Learning in Python, Pedregosa et al., JMLR 12, pp. 2825-2830, 2011."

    # Thank You
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "THANK YOU"
    content.text = ""

    prs.save('CyberGuard_Presentation.pptx')
    print("Presentation saved as CyberGuard_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
